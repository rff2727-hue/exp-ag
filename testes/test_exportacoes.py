"""Exportações: Excel, PDF, PNG, SVG, JSON e texto metodológico."""

import io
import json
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from aplicacao.exemplos import cenario_a, cenario_b, cenario_c_adicional
from motor.exportacao.croqui_estatico import exportar_croqui
from motor.exportacao.estilo import PREDEFINICOES
from motor.exportacao.excel import exportar_excel
from motor.exportacao.projeto_json import ErroImportacao, exportar_json, importar_json
from motor.exportacao.relatorio_pdf import relatorio_implantacao, relatorio_planejamento
from motor.exportacao.texto_metodologico import descricao_en, descricao_pt


@pytest.fixture(scope="module")
def plano():
    return cenario_a()


def test_excel_estrutura(plano):
    wb = load_workbook(io.BytesIO(exportar_excel(plano, ["Altura (cm)"])))
    assert wb.sheetnames == ["Resumo", "Tratamentos", "Casualização", "Croqui", "Planilha de campo",
                             "Coleta (formato longo)", "Quadro de GL", "Metodologia"]
    ws = wb["Casualização"]
    assert ws.max_row == 41
    codigos = [ws.cell(r, 5).value for r in range(2, 42)]
    assert codigos == [u.codigo for u in plano.casualizacao.unidades]
    cab = wb["Tratamentos"]["A1"]
    assert cab.font.b and cab.border.top.style == "medium" and cab.border.bottom.style == "thin"
    assert not wb["Tratamentos"]["A2"].font.b
    assert wb["Planilha de campo"].cell(1, 6).value == "Altura (cm)"
    assert wb["Croqui"].max_row >= 4


def test_pdfs(plano):
    for gerar in (relatorio_planejamento, relatorio_implantacao):
        dados = gerar(plano)
        assert dados[:4] == b"%PDF"
        leitor = PdfReader(io.BytesIO(dados))
        assert len(leitor.pages) >= 3
        texto = "".join(p.extract_text() for p in leitor.pages)
        assert plano.meta.codigo in texto


@pytest.mark.parametrize("fmt,assinatura", [("png", b"\x89PNG"), ("pdf", b"%PDF")])
def test_croqui_raster_e_pdf(plano, fmt, assinatura):
    for modo in ("estatistico", "operacional", "cientifico"):
        assert exportar_croqui(plano.casualizacao, plano.tratamentos, plano.croqui, PREDEFINICOES[modo], fmt,
                               "A4", 100)[:4] == assinatura


def test_croqui_svg_vetorial_com_texto(plano):
    svg = exportar_croqui(plano.casualizacao, plano.tratamentos, plano.croqui, PREDEFINICOES["cientifico"], "svg")
    raiz = ET.fromstring(svg)
    assert raiz.tag.endswith("svg")
    assert b"Bloco IV" in svg and b"T10" in svg  # texto editável, não convertido em curvas


def test_vasos_e_fatorial_exportam():
    for p in (cenario_b(), cenario_c_adicional()):
        assert exportar_excel(p)[:2] == b"PK"
        assert relatorio_planejamento(p)[:4] == b"%PDF"


def test_json_ida_e_volta(plano):
    p2 = importar_json(exportar_json(plano))
    assert [(u.codigo, u.tratamento) for u in p2.casualizacao.unidades] == \
        [(u.codigo, u.tratamento) for u in plano.casualizacao.unidades]
    assert p2.layout == plano.layout and p2.meta == plano.meta


def test_json_adulterado_e_detectado(plano):
    d = json.loads(exportar_json(plano))
    u = d["casualizacao"]["unidades"]
    u[0]["tratamento"], u[1]["tratamento"] = u[1]["tratamento"], u[0]["tratamento"]
    with pytest.raises(ErroImportacao):
        importar_json(json.dumps(d))
    with pytest.raises(ErroImportacao):
        importar_json(b"{nao e json")


def test_texto_metodologico_usa_apenas_o_configurado(plano):
    t = descricao_pt(plano)
    assert "blocos casualizados" in t and "dez tratamentos" in t and "quatro repetições" in t
    assert "40 unidades experimentais" in t and str(plano.casualizacao.semente) in t
    assert "cultivar" not in t  # não informado → não aparece
    assert "randomized complete block design" in descricao_en(plano)
    tb = descricao_pt(cenario_b())
    assert "inteiramente casualizado" in tb and "vasos" in tb and "mg kg⁻¹" in tb


def test_exemplos_json_do_repositorio_importam():
    from pathlib import Path

    arquivos = sorted((Path(__file__).resolve().parents[1] / "exemplos").glob("*.json"))
    assert len(arquivos) >= 3
    for arq in arquivos:
        p = importar_json(arq.read_bytes())
        assert p.verificar() == []
