"""Critérios de aceitação da Fase 1 (seção 18 da especificação)."""

import io

from openpyxl import load_workbook

from dados.banco import criar_engine, fabrica_sessoes, migrar
from dados.repositorios.experimentos import RepositorioExperimentos
from motor.dominio import EstruturaTratamentos, Fator, Nivel
from motor.exportacao.croqui_estatico import exportar_croqui
from motor.exportacao.estilo import PREDEFINICOES
from motor.exportacao.excel import exportar_excel
from motor.exportacao.relatorio_pdf import relatorio_planejamento
from motor.geometria.layout import ConfigLayout
from motor.planejamento import Metadados, planejar


def plano_a(semente=20261015):
    e = EstruturaTratamentos([Fator("Tratamento", [Nivel(f"T{i}") for i in range(1, 11)])])
    return planejar(Metadados("CEN-A", "Cenário A"), "DBC", e, 4, semente,
                    ConfigLayout(largura=3.0, comprimento=6.0, espaco_x=0.5, espaco_y=1.0, espaco_blocos=2.0,
                                 colunas_por_bloco=5, bordadura_lateral=0.5, bordadura_cabeceira=0.5),
                    "{exp}-P{parcela}-B{bloco}-T{trat}", "bloco")


def test_cenario_a_campo_dbc_10x4():
    p = plano_a()
    assert len(p.tratamentos) == 10 and p.casualizacao.n_unidades == 40
    assert p.verificar() == []
    for b in range(1, 5):
        assert sorted(u.tratamento for u in p.casualizacao.unidades if u.bloco == b) == list(range(1, 11))
    m = p.croqui.metricas
    assert m["n_unidades"] == 40 and m["area_parcela_m2"] == 18
    assert relatorio_planejamento(p)[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(exportar_excel(p)))
    assert wb["Casualização"].max_row == 41


def test_cenario_b_casa_de_vegetacao_dic_6x5():
    e = EstruturaTratamentos([Fator("Tratamento", [Nivel(f"T{i}") for i in range(1, 7)])])
    p = planejar(Metadados("CEN-B", "Cenário B", ambiente="casa de vegetação"), "DIC", e, 5, 7,
                 ConfigLayout(tipo="vasos", colunas_por_bloco=6), "{exp}-P{parcela}-T{trat}-R{rep}")
    us = p.casualizacao.unidades
    assert len(us) == 30 and len({u.codigo for u in us}) == 30
    assert all(u.codigo.startswith("CEN-B-P") for u in us)
    linhas = {c.linha for c in p.croqui.celulas}
    colunas = {c.coluna for c in p.croqui.celulas}
    assert linhas == set(range(5)) and colunas == set(range(6))
    for fmt, assin in (("png", b"\x89PNG"), ("svg", b"<?xm"), ("pdf", b"%PDF")):
        assert exportar_croqui(p.casualizacao, p.tratamentos, p.croqui, PREDEFINICOES["estatistico"],
                               fmt, dpi=100)[:4] == assin


def test_cenario_c_fatorial_3x2_r4():
    e = EstruturaTratamentos([Fator("A", [Nivel("a1"), Nivel("a2"), Nivel("a3")]),
                              Fator("B", [Nivel("b1"), Nivel("b2")])])
    for delin in ("DIC", "DBC"):
        p = planejar(Metadados("CEN-C", "Cenário C"), delin, e, 4, 11, ConfigLayout(), "{exp}-P{parcela}")
        combos = {(t.niveis["A"], t.niveis["B"]) for t in p.tratamentos}
        assert combos == {(a, b) for a in ("a1", "a2", "a3") for b in ("b1", "b2")}
        assert p.casualizacao.n_unidades == 24 and p.verificar() == []


def test_cenario_d_persistencia_apos_reinicio(url_temp):
    url = url_temp
    eng = criar_engine(url)
    migrar(eng)
    repo = RepositorioExperimentos(fabrica_sessoes(eng))
    original = plano_a()
    eid = repo.salvar_novo(repo.criar_projeto("PRJ", "Projeto"), original)
    eng.dispose()  # "encerra a sessão"
    del repo

    eng2 = criar_engine(url)  # "reabre o aplicativo"
    migrar(eng2)
    recuperado = RepositorioExperimentos(fabrica_sessoes(eng2)).carregar(eid)
    assert [(u.numero, u.parcela, u.codigo, u.tratamento, u.bloco) for u in recuperado.casualizacao.unidades] == \
        [(u.numero, u.parcela, u.codigo, u.tratamento, u.bloco) for u in original.casualizacao.unidades]
    assert [(t.codigo, t.rotulo) for t in recuperado.tratamentos] == [(t.codigo, t.rotulo) for t in original.tratamentos]
    assert [(c.numero, c.poligono) for c in recuperado.croqui.celulas] == \
        [(c.numero, c.poligono) for c in original.croqui.celulas]
    assert recuperado.croqui.metricas == original.croqui.metricas
    eng2.dispose()


def test_cenario_e_reprodutibilidade():
    a, b = plano_a(123456), plano_a(123456)
    assert [(u.codigo, u.tratamento) for u in a.casualizacao.unidades] == \
        [(u.codigo, u.tratamento) for u in b.casualizacao.unidades]
    assert [u.tratamento for u in plano_a(654321).casualizacao.unidades] != \
        [u.tratamento for u in a.casualizacao.unidades]
